# -*- coding: utf-8 -*-
from comun import *
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

TEAL, BANDA = "23656F", "E6F3FB"
F_CAB = Font(name="Lexend", size=9, bold=True, color="FFFFFF")
F_TXT = Font(name="Lexend", size=9, color="2B2B2B")
F_TOT = Font(name="Lexend", size=9, bold=True, color=TEAL)
R_CAB = PatternFill("solid", fgColor=TEAL); R_BAN = PatternFill("solid", fgColor=BANDA)
AL_I = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_IT = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center")
AL_D = Alignment(horizontal="right", vertical="center")
LINEA = Border(top=Side(style="thin", color=TEAL))
MON = '"$"#,##0.00'
HRS = 'General'
wb = Workbook()
CELDAS = {}   # key cells for the verifier

def cab(ws, fila, valores, anchos):
    for j, (v, w) in enumerate(zip(valores, anchos), start=1):
        c = ws.cell(row=fila, column=j, value=v); c.font = F_CAB; c.fill = R_CAB
        c.alignment = AL_C if j > 1 else AL_I
        if w: ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[fila].height = 22

def put(ws, r, j, v, font=F_TXT, fmt=None, al=None):
    c = ws.cell(row=r, column=j, value=v); c.font = font
    c.alignment = al or (AL_I if j == 1 else AL_D)
    if fmt: c.number_format = fmt
    return c

def banda(ws, r, ncol):
    if r % 2 == 0:
        for j in range(1, ncol + 1): ws.cell(row=r, column=j).fill = R_BAN

# 1. Parametros
ws = wb.active; ws.title = "Parámetros"
cab(ws, 1, ["Concepto", "Valor"], [46, 34])
par = [("Cliente", "Freshbox (%s)" % P["cliente"]), ("Atención", P["atencion"]),
       ("Fecha de emisión", P["fecha_emision"]),
       ("Vigencia de la tarifa ofertada", P["fecha_limite"]),
       ("Tarifa ofertada (MXN por hora)", P["tarifa_pref"]),
       ("Anticipo del esquema A", P["anticipo"]),
       ("Pagos mensuales del esquema B", C["pagos_b"]),
       ("Descuento del esquema C", P["descuento_unico"]),
       ("Plazo de ejecución", C["plazo"])]
if KEY == "integral":
    par += [("Horas del proyecto contable por separado", ruta.cifras("contabilidad")["tot"]),
            ("Horas del proyecto de inventario por separado", ruta.cifras("inventario")["tot"]),
            ("Pagos mensuales del esquema B por proyecto separado", ruta.cifras("contabilidad")["pagos_b"])]
for i, (k, v) in enumerate(par, start=2):
    put(ws, i, 1, k); put(ws, i, 2, v); banda(ws, i, 2)
ws.cell(row=6, column=2).number_format = MON
for r in (7, 9): ws.cell(row=r, column=2).number_format = "0%"
nota = {"contabilidad": "Sin desarrollos. Los documentos de compra y venta, con sus facturas, y la valoración se atienden en la propuesta de inventario.",
        "inventario": "Sin desarrollos. Los movimientos directos en contabilidad, la cobranza y la conciliación bancaria se atienden en la propuesta contable.",
        "integral": "Sin desarrollos. Reúne los proyectos contable y de inventario; las tareas comunes se ejecutan una sola vez."}[KEY]
PR, AN, NB, DS = ("Parámetros!$B$6", "Parámetros!$B$7", "Parámetros!$B$8", "Parámetros!$B$9")
NR = len(par) + 3
put(ws, NR, 1, nota); ws.merge_cells("A%d:B%d" % (NR, NR)); ws.row_dimensions[NR].height = 30

# 2. Ruta
ws = wb.create_sheet("Ruta")
INT = KEY == "integral"
cab(ws, 1, ["Núm.", "Aplicación", "Tarea", "Subtarea", "Tipo de trabajo", "Horas", "Hito", "Descripción"] + (["Proyecto"] if INT else []),
    [8, 20, 24, 32, 15, 8, 9, 80] + ([14] if INT else []))
NC = 9 if INT else 8
r = 2
for n, a, g, t, tw, hh, hi, de, pr in RUTA:
    for j, v in enumerate([n, a, g, t, tw, hh, hi, de] + ([pr] if INT else []), start=1):
        put(ws, r, j, v, al=AL_C if j in (1, 6, 7, 9) else AL_I)
    ws.cell(row=r, column=6).number_format = HRS
    banda(ws, r, NC)
    ws.row_dimensions[r].height = 15 * max(2, -(-len(de) // 95))
    r += 1
FIN = r - 1
put(ws, r, 4, "Total del alcance contratado", F_TOT, al=AL_D)
put(ws, r, 6, "=SUM(F2:F%d)" % FIN, F_TOT, HRS, al=AL_C)
for j in range(1, NC + 1): ws.cell(row=r, column=j).border = LINEA
CELDAS["ruta_total"] = ["Ruta", "F%d" % r]
ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:%s%d" % (L(NC), FIN)
RB, RE, RF, RG = ("Ruta!$B$2:$B$%d" % FIN, "Ruta!$E$2:$E$%d" % FIN,
                  "Ruta!$F$2:$F$%d" % FIN, "Ruta!$G$2:$G$%d" % FIN)

# 3. Resumen por etapa (por aplicacion)
ws = wb.create_sheet("Resumen por etapa")
cab(ws, 1, ["Aplicación", "Horas", "% del proyecto", "Importe"], [30, 10, 15, 20])
apps = []
for x in RUTA:
    if x[1] not in apps: apps.append(x[1])
TR = 2 + len(apps)
for i, a in enumerate(apps):
    r = 2 + i
    put(ws, r, 1, a); put(ws, r, 2, "=SUMIF(%s,$A%d,%s)" % (RB, r, RF), fmt=HRS)
    put(ws, r, 3, "=B%d/$B$%d" % (r, TR), fmt="0.0%")
    put(ws, r, 4, "=B%d*%s" % (r, PR), fmt=MON)
    banda(ws, r, 4)
    CELDAS["app " + a] = ["Resumen por etapa", "B%d" % r]
put(ws, TR, 1, "Total del alcance contratado", F_TOT)
for j, fm in ((2, HRS), (3, "0.0%"), (4, MON)):
    put(ws, TR, j, "=SUM(%s2:%s%d)" % (L(j), L(j), TR - 1), F_TOT, fm)
for j in range(1, 5): ws.cell(row=TR, column=j).border = LINEA
CELDAS.update(res_h=["Resumen por etapa", "B%d" % TR], res_pct=["Resumen por etapa", "C%d" % TR],
              res_p=["Resumen por etapa", "D%d" % TR])
ALC = "'Resumen por etapa'!$D$%d" % TR

# 4. Hitos (esquemas A, B y C)
ws = wb.create_sheet("Hitos")
cab(ws, 1, ["Esquema A — concepto", "Horas", "Importe"], [52, 10, 18])
put(ws, 2, 1, "Anticipo a la firma")
put(ws, 2, 3, "=ROUND(%s*%s,2)" % (ALC, AN), fmt=MON)
CELDAS["ant"] = ["Hitos", "C2"]
r = 3
for k, et in ETAPA.items():
    put(ws, r, 1, "%s. %s" % (k, et))
    put(ws, r, 2, '=SUMIF(%s,"%s",%s)' % (RG, k, RF), fmt=HRS)
    put(ws, r, 3, "=ROUND(B%d*%s*(1-%s),2)" % (r, PR, AN), fmt=MON)
    CELDAS[k + " h"] = ["Hitos", "B%d" % r]; CELDAS[k + " $"] = ["Hitos", "C%d" % r]
    banda(ws, r, 3)
    r += 1
ws.cell(row=r - 1, column=3).value = "=%s-SUM(C2:C%d)" % (ALC, r - 2)   # last milestone absorbs rounding
HR = r
put(ws, HR, 1, "Total antes de IVA", F_TOT)
put(ws, HR, 2, "=SUM(B3:B%d)" % (HR - 1), F_TOT, HRS)
put(ws, HR, 3, "=SUM(C2:C%d)" % (HR - 1), F_TOT, MON)
for j in range(1, 4): ws.cell(row=HR, column=j).border = LINEA
CELDAS["A total"] = ["Hitos", "C%d" % HR]
cab(ws, HR + 2, ["Esquemas B y C", "", "Importe"], [None, None, None])
put(ws, HR + 3, 1, "B. Pago mensual igual, al inicio de cada mes")
put(ws, HR + 3, 3, "=ROUND(%s/%s,2)" % (ALC, NB), fmt=MON)
put(ws, HR + 4, 1, "C. Descuento por pago único a la firma")
put(ws, HR + 4, 3, "=ROUND(%s*%s,2)" % (ALC, DS), fmt=MON)
put(ws, HR + 5, 1, "C. Pago único", F_TOT)
put(ws, HR + 5, 3, "=%s-C%d" % (ALC, HR + 4), F_TOT, MON)
put(ws, HR + 6, 1, "C. Tarifa efectiva por hora")
put(ws, HR + 6, 3, "=ROUND(%s*(1-%s),2)" % (PR, DS), fmt=MON)
CELDAS.update(B=["Hitos", "C%d" % (HR + 3)], Cdesc=["Hitos", "C%d" % (HR + 4)],
              Cpago=["Hitos", "C%d" % (HR + 5)], Ctarifa=["Hitos", "C%d" % (HR + 6)])
if INT:
    R0 = HR + 9
    cab(ws, R0, ["Comparativo de opciones", "Horas", "Importe"], [None, None, None])
    ws.cell(row=R0, column=4, value="Pago único (C)").font = F_CAB; ws.cell(row=R0, column=4).fill = R_CAB
    ws.cell(row=R0, column=5, value="Pago mensual (B)").font = F_CAB; ws.cell(row=R0, column=5).fill = R_CAB
    for cc in "DE": ws.column_dimensions[cc].width = 18
    ops = [("Contabilidad, cobranza y conciliación", "Parámetros!$B$11", "Parámetros!$B$13"),
           ("Inventario, compras, ventas y valoración", "Parámetros!$B$12", "Parámetros!$B$13"),
           ("Contabilidad e inventario", "'Resumen por etapa'!$B$%d" % TR, NB)]
    for i, (lab, hc, nb) in enumerate(ops):
        rr = R0 + 1 + i
        put(ws, rr, 1, lab); put(ws, rr, 2, "=" + hc, fmt=HRS)
        put(ws, rr, 3, "=B%d*%s" % (rr, PR), fmt=MON)
        put(ws, rr, 4, "=C%d-ROUND(C%d*%s,2)" % (rr, rr, DS), fmt=MON)
        put(ws, rr, 5, "=ROUND(C%d/%s,2)" % (rr, nb), fmt=MON)
        banda(ws, rr, 5)
        CELDAS["op%d h" % i] = ["Hitos", "B%d" % rr]; CELDAS["op%d $" % i] = ["Hitos", "C%d" % rr]
        CELDAS["op%d c" % i] = ["Hitos", "D%d" % rr]; CELDAS["op%d b" % i] = ["Hitos", "E%d" % rr]

# 5. Mezcla
ws = wb.create_sheet("Mezcla")
cab(ws, 1, ["Tipo de trabajo", "Horas", "% del proyecto"], [44, 10, 16])
tipos = sorted(set(x[4] for x in RUTA)); MR = 2 + len(tipos)
for i, t in enumerate(tipos):
    r = 2 + i
    put(ws, r, 1, t); put(ws, r, 2, "=SUMIF(%s,$A%d,%s)" % (RE, r, RF), fmt=HRS)
    put(ws, r, 3, "=B%d/$B$%d" % (r, MR), fmt="0.0%")
    banda(ws, r, 3)
put(ws, MR, 1, "Total del alcance contratado", F_TOT)
put(ws, MR, 2, "=SUM(B2:B%d)" % (MR - 1), F_TOT, HRS); put(ws, MR, 3, "=SUM(C2:C%d)" % (MR - 1), F_TOT, "0.0%")
for j in range(1, 4): ws.cell(row=MR, column=j).border = LINEA
put(ws, MR + 2, 1, "Configuración y datos")
put(ws, MR + 2, 2, '=SUMIF(%s,"Configuración",%s)+SUMIF(%s,"Datos",%s)' % (RE, RF, RE, RF), fmt=HRS)
put(ws, MR + 3, 1, "Levantamiento, conciliación, documentación, capacitación y cierre")
put(ws, MR + 3, 2, "=B%d-B%d" % (MR, MR + 2), fmt=HRS)
CELDAS.update(mez_total=["Mezcla", "B%d" % MR], mez_cd=["Mezcla", "B%d" % (MR + 2)])

wb.save(os.path.join(OUT, NX + ".xlsx"))
json.dump(CELDAS, open(os.path.join(OUT, "celdas.json"), "w"), ensure_ascii=False, indent=1)
print("ok", NX, "| FIN=%d TR=%d HR=%d MR=%d" % (FIN, TR, HR, MR))
