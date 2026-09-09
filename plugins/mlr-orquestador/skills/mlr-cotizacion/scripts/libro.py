# -*- coding: utf-8 -*-
"""Libro de captura del cuestionario de descubrimiento MLR."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import preguntas as Q

TEAL, GRIS, CLARO = "23656F", "595959", "E6F0F2"
OUT = "/mnt/user-data/outputs/Cuestionario de Descubrimiento Odoo - MLR.xlsx"

F_TIT = Font(name="Lexend", size=14, bold=True, color=TEAL)
F_CAB = Font(name="Lexend", size=10, bold=True, color="FFFFFF")
F_APP = Font(name="Lexend", size=11, bold=True, color="FFFFFF")
F_TXT = Font(name="Lexend", size=10)
F_NOTA = Font(name="Lexend", size=9, color=GRIS)
R_CAB = PatternFill("solid", fgColor=TEAL)
R_ALT = PatternFill("solid", fgColor=CLARO)
BORDE = Border(*[Side(style="thin", color="C9D8DC")] * 4)
AJUSTE = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ------------------------------------------------------- 1. Instrucciones
ws = wb.active; ws.title = "Instrucciones"
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 108
filas = [
    ("t", "Cuestionario de descubrimiento Odoo"),
    ("", ""),
    ("h", "Para qué sirve"),
    ("p", "Es el documento de trabajo de la reunión de descubrimiento. Se llena en vivo, "
          "delante del cliente, y de él sale la cotización. Lo que no quede aquí, no se cotiza."),
    ("", ""),
    ("h", "Cómo se usa en 30 minutos"),
    ("p", "1. La hoja «Ficha del sistema» se llena siempre, entre en el alcance lo que entre. "
          "Son cinco minutos y condiciona todo lo demás."),
    ("p", "2. En la hoja «Cuestionario», los bloques marcados NÚCLEO se preguntan siempre y en "
          "el orden en que están. Los demás solo se abren si el cliente los menciona."),
    ("p", "3. Cada bloque lleva sus minutos asignados. Si un bloque se pasa de tiempo, se anota "
          "«profundizar» y se sigue: el detalle fino se define en la etapa de descubrimiento pagada."),
    ("p", "4. La hoja «Migración» se llena al cierre, con la fecha de corte y quién entrega qué."),
    ("p", "5. La hoja «Resumen» se calcula sola y es lo que se lleva a la cotización."),
    ("", ""),
    ("h", "Si la reunión corre en Meet con Gemini tomando notas"),
    ("p", "Quien conduce dice en voz alta el nombre del bloque antes de empezarlo, para que la "
          "transcripción quede segmentada por aplicación."),
    ("p", "Las preguntas están redactadas para que la respuesta hablada se sostenga sola. Se leen "
          "como están: si se pregunta «¿en qué versión exacta están?», el cliente contesta «en la "
          "17.0», y esa línea sirve sin contexto. Si se pregunta «¿y eso?», la nota no sirve."),
    ("", ""),
    ("h", "Lo que NO se pregunta aquí"),
    ("p", "Nada que se pueda resolver auditando la base por API después de firmar el convenio de "
          "confidencialidad. Preguntar lo que uno mismo puede averiguar resta autoridad técnica y "
          "alarga la reunión. Este cuestionario es para lo que fija alcance y precio."),
    ("", ""),
    ("h", "Cómo se convierte en horas"),
] + [("l", d) for d in Q.DETONANTES]

r = 2
for tipo, txt in filas:
    c = ws.cell(row=r, column=2, value=txt)
    if tipo == "t": c.font = F_TIT
    elif tipo == "h": c.font = Font(name="Lexend", size=11, bold=True, color=TEAL)
    elif tipo == "l": c.font = F_TXT; c.value = "•  " + txt
    else: c.font = F_TXT
    c.alignment = AJUSTE
    ws.row_dimensions[r].height = 30 if tipo in ("p", "l") else 20
    r += 1

# --------------------------------------------------------- 2. Ficha
ws = wb.create_sheet("Ficha del sistema")
for col, anc in zip("ABCD", (4, 62, 34, 34)): ws.column_dimensions[col].width = anc
ws.cell(row=2, column=2, value="Ficha del sistema").font = F_TIT
ws.cell(row=3, column=2, value="Se llena siempre. Cinco minutos.").font = F_NOTA
for i, t in enumerate(["Pregunta", "Respuesta", "Notas"]):
    c = ws.cell(row=5, column=2 + i, value=t); c.font = F_CAB; c.fill = R_CAB
    c.alignment = Alignment(horizontal="center")
r = 6
for texto, tipo, ops in Q.FICHA:
    ws.cell(row=r, column=2, value=texto).font = F_TXT
    ws.cell(row=r, column=2).alignment = AJUSTE
    for col in (2, 3, 4):
        ws.cell(row=r, column=col).border = BORDE
        if r % 2 == 0: ws.cell(row=r, column=col).fill = R_ALT
    ws.cell(row=r, column=3).font = F_TXT
    ws.cell(row=r, column=4).font = F_TXT
    if tipo == "opcion" and ops:
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(ops), allow_blank=True)
        ws.add_data_validation(dv); dv.add(ws.cell(row=r, column=3))
    elif tipo == "si/no":
        dv = DataValidation(type="list", formula1='"Sí,No,Por definir"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(ws.cell(row=r, column=3))
    ws.row_dimensions[r].height = 26
    r += 1

# ----------------------------------------------------- 3. Cuestionario
ws = wb.create_sheet("Cuestionario")
anchos = [4, 10, 58, 26, 14, 16, 12, 30]
for i, a in enumerate(anchos): ws.column_dimensions[get_column_letter(i + 1)].width = a
ws.cell(row=2, column=2, value="Cuestionario por aplicación").font = F_TIT
ws.cell(row=3, column=2, value="Los bloques NÚCLEO se preguntan siempre y en este orden. "
        "Los demás solo se abren si el cliente los menciona.").font = F_NOTA
cabs = ["¿Aplica?", "Pregunta", "Respuesta", "Volumen", "¿Desarrollo?", "Horas", "Notas"]
for i, t in enumerate(cabs):
    c = ws.cell(row=5, column=2 + i, value=t); c.font = F_CAB; c.fill = R_CAB
    c.alignment = Alignment(horizontal="center", wrap_text=True)

dv_ap = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
dv_sn = DataValidation(type="list", formula1='"Sí,No,Por definir"', allow_blank=True)
dv_de = DataValidation(type="list", formula1='"No,Sí — a estimar,Sí — estimado"', allow_blank=True)
for dv in (dv_ap, dv_sn, dv_de): ws.add_data_validation(dv)

r = 6
indice = []
for clave, nombre, minutos, nucleo, preguntas in Q.BLOQUES:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    etq = "%s   ·   %d min   ·   %s" % (nombre, minutos, "NÚCLEO" if nucleo else "condicional")
    c = ws.cell(row=r, column=2, value=etq)
    c.font = F_APP; c.fill = R_CAB; c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 24
    indice.append((nombre, minutos, nucleo, r))
    r += 1
    for texto, tipo, ops in preguntas:
        ws.cell(row=r, column=3, value=texto).font = F_TXT
        ws.cell(row=r, column=3).alignment = AJUSTE
        for col in range(2, 9):
            ws.cell(row=r, column=col).border = BORDE
            ws.cell(row=r, column=col).font = F_TXT
        dv_ap.add(ws.cell(row=r, column=2))
        dv_de.add(ws.cell(row=r, column=6))
        if tipo == "opcion" and ops:
            d = DataValidation(type="list", formula1='"%s"' % ",".join(ops), allow_blank=True)
            ws.add_data_validation(d); d.add(ws.cell(row=r, column=4))
        elif tipo == "si/no":
            dv_sn.add(ws.cell(row=r, column=4))
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1

# bloque de cierre
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
c = ws.cell(row=r, column=2, value="Cierre de la reunión   ·   3 min   ·   NÚCLEO")
c.font = F_APP; c.fill = R_CAB; c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[r].height = 24
r += 1
for texto, tipo, ops in Q.CIERRE:
    ws.cell(row=r, column=3, value=texto).font = F_TXT
    ws.cell(row=r, column=3).alignment = AJUSTE
    for col in range(2, 9):
        ws.cell(row=r, column=col).border = BORDE
        ws.cell(row=r, column=col).font = F_TXT
    dv_ap.add(ws.cell(row=r, column=2))
    ws.row_dimensions[r].height = 26
    r += 1
ULT = r - 1
ws.freeze_panes = "C6"
ws.auto_filter.ref = "B5:H%d" % ULT

# ------------------------------------------------------- 4. Migración
ws = wb.create_sheet("Migración")
for col, anc in zip("ABCDEFG", (4, 34, 26, 18, 22, 18, 30)): ws.column_dimensions[col].width = anc
ws.cell(row=2, column=2, value="Datos a migrar").font = F_TIT
ws.cell(row=3, column=2, value="Se llena al cierre. Una línea por catálogo o saldo.").font = F_NOTA
for i, t in enumerate(["Qué se migra", "Desde dónde", "Formato", "Volumen aprox.",
                       "Fecha de corte", "Quién lo entrega"]):
    c = ws.cell(row=5, column=2 + i, value=t); c.font = F_CAB; c.fill = R_CAB
    c.alignment = Alignment(horizontal="center", wrap_text=True)
base = ["Catálogo de productos", "Clientes", "Proveedores", "Listas de precio",
        "Saldos de clientes por cobrar", "Saldos de proveedores por pagar",
        "Existencias y costos de inventario", "Catálogo de cuentas",
        "Balanza de comprobación al corte", "Activos fijos", "Obras o proyectos en curso"]
for i, n in enumerate(base):
    rr = 6 + i
    ws.cell(row=rr, column=2, value=n).font = F_TXT
    for col in range(2, 8):
        ws.cell(row=rr, column=col).border = BORDE
        ws.cell(row=rr, column=col).font = F_TXT
        if rr % 2 == 0: ws.cell(row=rr, column=col).fill = R_ALT
    ws.row_dimensions[rr].height = 22

# --------------------------------------------------------- 5. Resumen
ws = wb.create_sheet("Resumen")
for col, anc in zip("ABC", (4, 62, 24)): ws.column_dimensions[col].width = anc
ws.cell(row=2, column=2, value="Resumen para cotizar").font = F_TIT
ws.cell(row=3, column=2, value="Se calcula solo desde la hoja «Cuestionario».").font = F_NOTA
res = [
 ("Preguntas marcadas como aplicables",
  '=COUNTIF(Cuestionario!B:B,"Sí")'),
 ("Puntos que salieron como desarrollo",
  '=COUNTIF(Cuestionario!F:F,"Sí — a estimar")+COUNTIF(Cuestionario!F:F,"Sí — estimado")'),
 ("Horas capturadas en la reunión",
  '=SUM(Cuestionario!G:G)'),
 ("Procesos que hoy no existen",
  '=COUNTIF(Cuestionario!D:D,"No existe")'),
 ("Procesos que están en Odoo a medias",
  '=COUNTIF(Cuestionario!D:D,"En Odoo, a medias")'),
 ("Procesos que hoy viven en Excel o en papel",
  '=COUNTIF(Cuestionario!D:D,"En Excel o en papel")'),
 ("Catálogos a migrar con histórico",
  '=COUNTIF(Cuestionario!D:D,"Catalogo con histórico")'),
]
for i, (n, f) in enumerate(res):
    rr = 5 + i
    ws.cell(row=rr, column=2, value=n).font = F_TXT
    c = ws.cell(row=rr, column=3, value=f)
    c.font = Font(name="Lexend", size=11, bold=True, color=TEAL)
    c.alignment = Alignment(horizontal="center")
    for col in (2, 3):
        ws.cell(row=rr, column=col).border = BORDE
    ws.row_dimensions[rr].height = 24
rr = 5 + len(res) + 2
ws.cell(row=rr, column=2, value="Cómo se lee este resumen").font = Font(
    name="Lexend", size=11, bold=True, color=TEAL)
for i, d in enumerate(Q.DETONANTES):
    c = ws.cell(row=rr + 1 + i, column=2, value="•  " + d)
    c.font = F_TXT; c.alignment = AJUSTE
    ws.row_dimensions[rr + 1 + i].height = 30

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
tot = sum(m for _, _, m, n, _ in Q.BLOQUES if n) + 5 + 3
print("OK", OUT)
print("bloques: %d (núcleo %d) | preguntas: %d | núcleo en minutos: %d"
      % (len(Q.BLOQUES), sum(1 for b in Q.BLOQUES if b[3]),
         len(Q.FICHA) + sum(len(b[4]) for b in Q.BLOQUES) + len(Q.CIERRE), tot))
