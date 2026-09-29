# -*- coding: utf-8 -*-
"""Closing controls of one journal and period. Reads Odoo (read only), writes the "lo escribe la skill"
values into the Verificación sheet, recalculates and says CERRADO or ABIERTO with the reasons.

    python3 verificar.py <params.json> <libro.xlsx> <carpeta_trabajo> [--sin-odoo]

If the analyst registers something by hand in Odoo with the same API user during the work, it shows up
in "creados sin aprobación": explain it in the chat and record it as a resolved action.
Controls: PDF read against the cover, no unidentified items, items to review only with an accepted
explanation, cash flow explained to zero, Odoo closing balance equal to the bank, approved actions all
applied and verified, statement lines left unreconciled, suspense balance, payments and entries created
without approval (Odoo count against the log), and the trial balance adding up to 0.00.
It also prints the six questions of the Resumen so the chat can report them before opening Excel.
"""
import argparse
import collections
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from llenar_libro import recalcula  # noqa: E402


def lee_odoo(P, carpeta):
    from odoo_client import Lectura
    r = Lectura.desde_entorno(compania_id=P["compania_id"])
    cia = [("company_id", "=", P["compania_id"])]
    dia = r.call("account.journal", "read", [P["diario_id"]], fields=["suspense_account_id"])[0]
    sin = r.cnt("account.bank.statement.line", [("journal_id", "=", P["diario_id"]), ("is_reconciled", "=", False), ("date", "<=", P["hasta"])])
    susp = 0.0
    if dia.get("suspense_account_id"):
        susp = round(sum(x["balance"] for x in r.sr_todo("account.move.line", [("account_id", "=", dia["suspense_account_id"][0]),
                                                                              ("parent_state", "=", "posted"), ("date", "<=", P["hasta"])] + cia, ["balance"])), 2)
    balanza = round(sum(x["balance"] for x in r.sr_todo("account.move.line", [("parent_state", "=", "posted"), ("date", "<=", P["hasta"])] + cia, ["balance"])), 2)
    creados = None
    ses = os.path.join(carpeta, "sesion.json")
    if os.path.exists(ses):
        inicio = json.load(open(ses))["inicio_utc"]      # UTC, same clock as Odoo's create_date
        lineas = [json.loads(x) for x in open(os.path.join(carpeta, "bitacora.jsonl")) if x.strip()] if os.path.exists(os.path.join(carpeta, "bitacora.jsonl")) else []
        aprob = collections.Counter()
        for l in lineas:
            if str(l.get("resultado", "")).startswith("ERROR"):
                continue
            if l["metodo"] == "action_create_payments":
                aprob["pago"] += 1
            if l["metodo"] == "create" and l["modelo"] == "account.move":
                aprob["asiento"] += 1
        pagos = r.cnt("account.payment", [("create_uid", "=", r.uid), ("create_date", ">=", inicio)] + cia)
        asientos = r.cnt("account.move", [("create_uid", "=", r.uid), ("create_date", ">=", inicio), ("move_type", "=", "entry"),
                                          ("origin_payment_id", "=", False), ("tax_cash_basis_origin_move_id", "=", False),
                                          ("statement_line_id", "=", False)] + cia)
        creados = max(0, pagos - aprob["pago"]) + max(0, asientos - aprob["asiento"])
    # None = unknown (no session start recorded): the control stays Pendiente instead of a false OK
    return {"sin_conciliar": sin, "suspenso": susp, "balanza": balanza, "creados_sin_aprobacion": creados}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("params")
    ap.add_argument("libro")
    ap.add_argument("carpeta")
    ap.add_argument("--sin-odoo", action="store_true")
    a = ap.parse_args()
    P = json.load(open(a.params))
    wb = openpyxl.load_workbook(a.libro)
    ws = wb["Verificación"]
    if not a.sin_odoo:
        o = lee_odoo(P, a.carpeta)
        for r in range(8, ws.max_row + 1):
            n = ws.cell(r, 1).value or ""
            if n.startswith("Líneas de extracto"):
                ws.cell(r, 2).value = o["sin_conciliar"]
            elif n.startswith("Saldo de la cuenta de suspenso"):
                ws.cell(r, 2).value = o["suspenso"]
            elif n.startswith("Pagos y asientos creados"):
                ws.cell(r, 2).value = o["creados_sin_aprobacion"]
            elif n.startswith("Suma total de la balanza"):
                ws.cell(r, 2).value = o["balanza"]
        wb.save(a.libro)
    res = recalcula(a.libro)
    if res[0] is None:
        sys.exit(res[1][0])
    n, err, calc = res
    if err:
        print("Fórmulas con error:", err[:10])
    v = openpyxl.load_workbook(calc, data_only=True)
    vs, rs = v["Verificación"], v["Resumen"]
    malos = []
    for r in range(8, vs.max_row + 1):
        if vs.cell(r, 4).value in ("Revisar", "Pendiente"):
            malos.append("%s: %s (esperado %s)" % (vs.cell(r, 1).value, vs.cell(r, 2).value, vs.cell(r, 3).value))
        if vs.cell(r, 1).value == "Estado del diario":
            estado = vs.cell(r, 4).value
    print("1. PDF: %s" % rs["C8"].value)
    print("2. Movimientos: %s (conciliado %s, validar %s, revisar %s, no identificado %s de %s)" % (
        rs["G8"].value, rs["H9"].value, rs["H10"].value, rs["H11"].value, rs["H12"].value, rs["H13"].value))
    print("3. Flujo explicado: diferencia cobros %s, pagos %s" % (rs["D23"].value, rs["D24"].value))
    print("4. Alertas: %s" % rs["G18"].value)
    print("5. %s · %s" % (rs["C28"].value, rs["C35"].value))
    print("6. %s IVA trasladado Odoo - conciliación %s; acreditable %s" % (rs["G28"].value, rs["H31"].value, rs["H32"].value))
    print("Diario %s" % estado)
    for m in malos:
        print("  -", m)
    os.replace(calc, a.libro.replace(".xlsx", "") + ".calc.xlsx")
    sys.exit(0 if estado == "CERRADO" else 4)


if __name__ == "__main__":
    main()
