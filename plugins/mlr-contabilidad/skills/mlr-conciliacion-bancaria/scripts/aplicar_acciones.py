# -*- coding: utf-8 -*-
"""Apply ONLY the approved actions of the workbook to Odoo, one type at a time, with revalidation,
JSON backup before each write, read-back verification and log.

    python3 aplicar_acciones.py <params.json> <libro.xlsx> <carpeta_trabajo>            # dry run (default)
    python3 aplicar_acciones.py ... --aplicar                                            # first pending of each new type
    python3 aplicar_acciones.py ... --aplicar --lote                                     # the rest, only for types already validated
    python3 aplicar_acciones.py ... --validar-tipo registrar_pago                        # after the browser check of the first one
    add --confirmo-produccion when params["entorno"] == "produccion"

Rules enforced here, not only in the instructions:
  - Aprobado = "Sí" in the Acciones sheet and a complete operation JSON; "Modificar" is never applied.
  - Revalidation against Odoo right before writing; if the document is no longer open or the amount
    changed, the row goes to "Revalidar" and nothing is written.
  - Test before batch: a type of operation runs on a single row until the user confirms, after
    checking it in the browser, that it looks right (--validar-tipo). Only then --lote applies the rest.
  - Forbidden methods (unlink, unreconcile, lock dates, reversals) are refused by odoo_client.Escritura.
  - Methods that return None are verified by reading, never retried blindly.
Close Excel before running: the Estado, IDs resultado and Verificado columns are written back.
"""
import argparse
import base64
import datetime
import json
import os
import sys

import openpyxl

from odoo_client import Escritura, Lectura

COL = {"id": 1, "tipo": 2, "partida": 3, "importe": 4, "operacion": 10, "aprobado": 11, "comentario": 12, "estado": 13, "ids": 14, "verificado": 15}
PERMITIDOS = {  # (model, field) pairs an "escribir_campos" action may touch
    "res.partner": {"vat", "zip", "l10n_mx_edi_fiscal_regime", "country_id"},
    "res.company": {"tax_cash_basis_journal_id"},
    "account.tax": {"cash_basis_transition_account_id"},
}
TOL = 0.01


class Revalidar(Exception):
    pass


def cuenta_id(r, codigo, cia):
    res = r.sr("account.account", [("code", "=", codigo)], ["id", "active"], context={"allowed_company_ids": [cia]})
    if not res:
        raise Revalidar("No existe la cuenta %s" % codigo)
    return res[0]["id"]


def no_bloqueada(r, fecha, cia):
    c = r.call("res.company", "read", [cia], fields=["fiscalyear_lock_date", "tax_lock_date", "hard_lock_date"])[0]
    for k, v in c.items():
        if k != "id" and v and fecha <= v:
            raise Revalidar("La fecha %s cae en un periodo bloqueado (%s = %s)" % (fecha, k, v))


# ---------------------------------------------------------------- operations: revalidate -> write -> verify
def op_registrar_pago(r, w, op, P, simula):
    facts = r.call("account.move", "read", op["factura_ids"], fields=["state", "amount_residual", "partner_id", "move_type", "payment_state", "name"])
    for f in facts:
        if f["state"] != "posted" or f["payment_state"] in ("paid", "reversed"):
            raise Revalidar("%s ya no está abierta (%s / %s)" % (f["name"], f["state"], f["payment_state"]))
    resid = round(sum(f["amount_residual"] for f in facts), 2)
    if op["importe"] > resid + TOL:
        raise Revalidar("El importe %.2f es mayor que el saldo %.2f de %s" % (op["importe"], resid, ", ".join(f["name"] for f in facts)))
    if len({f["partner_id"][0] for f in facts}) > 1:
        raise Revalidar("Las facturas son de contactos distintos")
    no_bloqueada(r, op["fecha"], P["compania_id"])
    if simula:
        return "Registraría un pago de %.2f el %s contra %s" % (op["importe"], op["fecha"], ", ".join(f["name"] for f in facts)), None
    ctx = {"active_model": "account.move", "active_ids": op["factura_ids"]}
    vals = {"payment_date": op["fecha"], "amount": op["importe"], "journal_id": op.get("diario_id") or P["diario_id"],
            "communication": op.get("memo") or "", "group_payment": True}
    if op.get("forma_pago"):
        fp = r.sr("l10n_mx_edi.payment.method", [("code", "=", op["forma_pago"])], ["id"])
        if fp:
            vals["l10n_mx_edi_payment_method_id"] = fp[0]["id"]
    antes = {p["id"] for p in r.sr("account.payment", [("partner_id", "=", facts[0]["partner_id"][0]), ("date", "=", op["fecha"])], ["id"])}
    wiz = w.crea("account.payment.register", vals, set(vals), contexto=ctx)
    w.ejecuta("account.payment.register", "action_create_payments", [[wiz]], {"context": ctx}, nuevos={"wizard": wiz})
    despues = r.sr("account.payment", [("partner_id", "=", facts[0]["partner_id"][0]), ("date", "=", op["fecha"])], ["id", "name", "amount", "state"])
    nuevos = [p for p in despues if p["id"] not in antes]
    facts2 = r.call("account.move", "read", op["factura_ids"], fields=["amount_residual"])
    bajo = round(resid - sum(f["amount_residual"] for f in facts2), 2)
    ok = len(nuevos) == 1 and abs(bajo - op["importe"]) <= TOL
    return ("Pago %s por %.2f; saldo de las facturas bajó %.2f" % (nuevos[0]["name"] if nuevos else "?", op["importe"], bajo),
            {"ids": [p["id"] for p in nuevos], "verificado": ok})


def op_asiento_manual(r, w, op, P, simula):
    cia = P["compania_id"]
    lineas = []
    for l in op["lineas"]:
        lineas.append({"account_id": cuenta_id(r, l["cuenta"], cia), "debit": round(l.get("debe") or 0, 2), "credit": round(l.get("haber") or 0, 2),
                       "name": l.get("nombre") or op.get("referencia", ""), **({"partner_id": l["partner_id"]} if l.get("partner_id") else {})})
    dif = round(sum(x["debit"] - x["credit"] for x in lineas), 2)
    if dif:
        raise Revalidar("El asiento no cuadra por %.2f" % dif)
    no_bloqueada(r, op["fecha"], cia)
    if simula:
        return "Crearía y publicaría un asiento de %.2f el %s en el diario %s" % (sum(x["debit"] for x in lineas), op["fecha"], op.get("diario_id")), None
    vals = {"move_type": "entry", "journal_id": op["diario_id"], "date": op["fecha"], "ref": op.get("referencia", ""),
            "line_ids": [(0, 0, x) for x in lineas]}
    mid = w.crea("account.move", vals, set(vals))
    w.ejecuta("account.move", "action_post", [[mid]], ids=[mid], campos_respaldo=["state"], nuevos={"state": "posted"})
    m = r.call("account.move", "read", [mid], fields=["name", "state", "amount_total"])[0]
    return "Asiento %s %s por %.2f" % (m["name"], m["state"], m["amount_total"]), {"ids": [mid], "verificado": m["state"] == "posted"}


def op_linea_extracto(r, w, op, P, simula):
    """Reconcile a statement line parked in suspense against accounts (technique proven in saas~19.4)."""
    sl = r.call("account.bank.statement.line", "read", [op["statement_line_id"]], fields=["is_reconciled", "move_id", "amount", "journal_id", "date"])[0]
    if sl["is_reconciled"]:
        raise Revalidar("La línea de extracto ya está conciliada")
    no_bloqueada(r, sl["date"], P["compania_id"])
    susp = r.call("account.journal", "read", [sl["journal_id"][0]], fields=["suspense_account_id"])[0]["suspense_account_id"][0]
    mid = sl["move_id"][0]
    ls = r.sr("account.move.line", [("move_id", "=", mid), ("account_id", "=", susp)], ["id", "balance"])
    if len(ls) != 1:
        raise Revalidar("La línea no tiene exactamente un apunte en suspenso (%d)" % len(ls))
    total = round(-ls[0]["balance"], 2)
    contra = []
    for c in op["contrapartidas"]:
        imp = round(c["importe"], 2)
        contra.append({"account_id": cuenta_id(r, c["cuenta"], P["compania_id"]), "debit": imp if total < 0 else 0, "credit": imp if total > 0 else 0,
                       "name": c.get("nombre") or "", **({"partner_id": c["partner_id"]} if c.get("partner_id") else {})})
    if abs(sum(c["importe"] for c in op["contrapartidas"]) - abs(total)) > 0.005:
        raise Revalidar("Las contrapartidas suman %.2f y el suspenso es %.2f" % (sum(c["importe"] for c in op["contrapartidas"]), abs(total)))
    if simula:
        return "Cambiaría el suspenso de %.2f por %d contrapartida(s) y volvería a publicar" % (abs(total), len(contra)), None
    ctx = {"skip_account_move_synchronization": True}
    # full copy of the suspense line, so the change can be undone from the backup
    w._respaldo("account.move.line", [ls[0]["id"]], {"account_id", "partner_id", "debit", "credit", "balance", "amount_currency",
                                                     "currency_id", "name", "date", "move_id"})
    w.ejecuta("account.move", "button_draft", [[mid]], ids=[mid], campos_respaldo=["state", "line_ids"])
    try:
        w.escribe("account.move", [mid], {"line_ids": [(2, ls[0]["id"])] + [(0, 0, c) for c in contra]}, {"line_ids"}, contexto=ctx)
    finally:
        w.ejecuta("account.move", "action_post", [[mid]], ids=[mid], nuevos={"state": "posted"})
    ok = r.call("account.bank.statement.line", "read", [op["statement_line_id"]], fields=["is_reconciled"])[0]["is_reconciled"]
    msg = "Línea de extracto %s conciliada: %s" % (op["statement_line_id"], ok)
    if ok and op.get("conciliar_con_factura_id"):
        # the new receivable/payable line of the statement move against the invoice's open line on the same account
        cta = contra[0]["account_id"]
        nueva = r.sr("account.move.line", [("move_id", "=", mid), ("account_id", "=", cta), ("reconciled", "=", False)], ["id"])
        abierta = r.sr("account.move.line", [("move_id", "=", op["conciliar_con_factura_id"]), ("account_id", "=", cta), ("reconciled", "=", False)], ["id"])
        if nueva and abierta:
            ids = [nueva[0]["id"], abierta[0]["id"]]
            w.ejecuta("account.move.line", "reconcile", [ids], ids=ids, campos_respaldo=["reconciled", "amount_residual"])
            ok = r.call("account.move.line", "read", [nueva[0]["id"]], fields=["reconciled"])[0]["reconciled"]
            msg += "; aplicada a la factura: %s" % ok
        else:
            ok = False
            msg += "; no encontré la línea abierta de la factura en la cuenta %s" % cta
    return msg, {"ids": [mid], "verificado": ok}


def op_conciliar_apuntes(r, w, op, P, simula):
    ls = r.call("account.move.line", "read", op["line_ids"], fields=["reconciled", "account_id", "partner_id", "amount_residual"])
    if any(l["reconciled"] for l in ls):
        raise Revalidar("Algún apunte ya está conciliado")
    if len({l["account_id"][0] for l in ls}) != 1:
        raise Revalidar("Los apuntes son de cuentas distintas")
    if simula:
        return "Conciliaría %d apuntes con saldo neto %.2f" % (len(ls), sum(l["amount_residual"] for l in ls)), None
    w.ejecuta("account.move.line", "reconcile", [op["line_ids"]], ids=op["line_ids"], campos_respaldo=["reconciled", "amount_residual"])
    ls2 = r.call("account.move.line", "read", op["line_ids"], fields=["reconciled", "amount_residual"])  # reconcile() returns None: read to verify
    ok = any(l["reconciled"] for l in ls2) or abs(sum(l["amount_residual"] for l in ls2)) < abs(sum(l["amount_residual"] for l in ls))
    return "Apuntes conciliados: %s" % ok, {"ids": op["line_ids"], "verificado": ok}


def op_ligar_xml(r, w, op, P, simula):
    f = r.call("account.move", "read", [op["factura_id"]], fields=["name", "l10n_mx_edi_cfdi_uuid", "move_type", "state"])[0]
    if f.get("l10n_mx_edi_cfdi_uuid"):
        raise Revalidar("%s ya tiene folio fiscal %s" % (f["name"], f["l10n_mx_edi_cfdi_uuid"]))
    raw = open(op["xml_ruta"], "rb").read()
    if simula:
        return "Ligaría %s (%d bytes) a %s" % (os.path.basename(op["xml_ruta"]), len(raw), f["name"]), None
    att = w.crea("ir.attachment", {"name": os.path.basename(op["xml_ruta"]), "res_model": "account.move", "res_id": f["id"],
                                   "raw": base64.b64encode(raw).decode(), "mimetype": "application/xml"},
                 {"name", "res_model", "res_id", "raw", "mimetype"})  # saas~19.4 has no 'datas' field
    estado = "invoice_received" if f["move_type"].startswith("in_") else "invoice_sent"
    doc = w.crea("l10n_mx_edi.document", {"move_id": f["id"], "invoice_ids": [(6, 0, [f["id"]])], "state": estado, "sat_state": "skip",
                                          "datetime": op.get("fecha_timbrado") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
                 {"move_id", "invoice_ids", "state", "sat_state", "datetime"})
    w.escribe("l10n_mx_edi.document", [doc], {"attachment_id": att, "sat_state": "not_defined"}, {"attachment_id", "sat_state"})
    try:
        w.ejecuta("account.move", "l10n_mx_edi_cfdi_try_sat", [[f["id"]]], ids=[f["id"]])
    except Exception as e:  # the SAT query can fail on its own; the link is already made
        print("   aviso: consulta al SAT: %s" % e)
    u = r.call("account.move", "read", [f["id"]], fields=["l10n_mx_edi_cfdi_uuid"])[0]["l10n_mx_edi_cfdi_uuid"]
    return "%s con folio fiscal %s" % (f["name"], u or "(no apareció: quitar y reponer el adjunto)"), {"ids": [att, doc], "verificado": bool(u)}


def op_adjuntar_rep(r, w, op, P, simula):
    raw = open(op["xml_ruta"], "rb").read()
    if simula:
        return "Adjuntaría el REP %s a %s %s con nota" % (os.path.basename(op["xml_ruta"]), op["res_model"], op["res_id"]), None
    att = w.crea("ir.attachment", {"name": os.path.basename(op["xml_ruta"]), "res_model": op["res_model"], "res_id": op["res_id"],
                                   "raw": base64.b64encode(raw).decode(), "mimetype": "application/xml"},
                 {"name", "res_model", "res_id", "raw", "mimetype"})
    # plain text body: message_post over RPC escapes HTML
    w.ejecuta(op["res_model"], "message_post", [[op["res_id"]]], {"body": op["nota"], "attachment_ids": [att]}, ids=[op["res_id"]])
    n = r.cnt("ir.attachment", [("id", "=", att)])
    return "REP adjunto (%d)" % att, {"ids": [att], "verificado": n == 1}


def op_escribir_campos(r, w, op, P, simula):
    modelo, vals = op["modelo"], op["valores"]
    fuera = set(vals) - PERMITIDOS.get(modelo, set())
    if fuera:
        raise Revalidar("Campos fuera de la lista permitida para %s: %s" % (modelo, sorted(fuera)))
    if modelo == "res.partner" and "country_id" in vals and "l10n_mx_edi_fiscal_regime" not in vals:
        raise Revalidar("Al poner país, Odoo asigna régimen 601 solo: incluye el régimen correcto en el mismo cambio")
    if modelo == "account.tax":
        abiertas = r.cnt("account.move.line", [("tax_ids", "in", op["ids"]), ("move_id.payment_state", "in", ["not_paid", "partial"]),
                                               ("parent_state", "=", "posted")])
        if abiertas:
            raise Revalidar("Hay %d apuntes de facturas abiertas con ese impuesto: cambiar su cuenta de tránsito descuadra el flujo" % abiertas)
    antes = r.call(modelo, "read", op["ids"], fields=list(vals))
    if simula:
        return "Cambiaría %s en %s %s (antes: %s)" % (vals, modelo, op["ids"], antes), None
    w.escribe(modelo, op["ids"], vals, set(vals))
    desp = r.call(modelo, "read", op["ids"], fields=list(vals))
    ok = all((x[k][0] if isinstance(x[k], list) else x[k]) == v for x in desp for k, v in vals.items())
    return "Actualizado %s %s" % (modelo, op["ids"]), {"ids": op["ids"], "verificado": ok}


OPS = {"registrar_pago": op_registrar_pago, "asiento_manual": op_asiento_manual, "linea_extracto": op_linea_extracto,
       "conciliar_apuntes": op_conciliar_apuntes, "ligar_xml": op_ligar_xml, "adjuntar_rep": op_adjuntar_rep,
       "escribir_campos": op_escribir_campos}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("params")
    ap.add_argument("libro")
    ap.add_argument("carpeta")
    ap.add_argument("--aplicar", action="store_true")
    ap.add_argument("--lote", action="store_true")
    ap.add_argument("--confirmo-produccion", action="store_true")
    ap.add_argument("--validar-tipo")
    a = ap.parse_args()
    P = json.load(open(a.params))
    val_path = os.path.join(a.carpeta, "tipos_validados.json")
    validados = set(json.load(open(val_path))) if os.path.exists(val_path) else set()
    if a.validar_tipo:
        validados.add(a.validar_tipo)
        json.dump(sorted(validados), open(val_path, "w"))
        pp = os.path.join(a.carpeta, "tipos_en_prueba.json")
        if os.path.exists(pp):
            json.dump(sorted(set(json.load(open(pp))) - {a.validar_tipo}), open(pp, "w"))
        print("Tipo validado para lote:", a.validar_tipo)
        return
    if a.aplicar:
        if P.get("entorno") not in ("pruebas", "produccion"):
            raise SystemExit("params.json no declara el entorno (pruebas o produccion): no se escribe nada hasta que la persona lo diga.")
        if P["entorno"] == "produccion" and P.get("respaldo_base") != datetime.date.today().isoformat():
            raise SystemExit("Base de producción: falta en params.json \"respaldo_base\" con la fecha de hoy, el respaldo de la base que la persona tomó hoy.")
    prueba_path = os.path.join(a.carpeta, "tipos_en_prueba.json")
    en_prueba = set(json.load(open(prueba_path))) if os.path.exists(prueba_path) else set()
    wb = openpyxl.load_workbook(a.libro)
    ws = wb["Acciones"]
    r = Lectura.desde_entorno(compania_id=P["compania_id"])
    w = Escritura(r, a.carpeta, "sesion-%s" % datetime.date.today(), P["entorno"], a.confirmo_produccion) if a.aplicar else None
    hechos_tipo, sin_op, pendientes = set(), [], 0
    for fila in range(8, ws.max_row + 1):
        g = lambda k: ws.cell(fila, COL[k]).value
        if not g("id") or g("aprobado") != "Sí" or g("estado") not in (None, "", "Pendiente", "Revalidar"):
            if g("aprobado") == "Modificar":
                print("%s marcada Modificar: se ajusta en el chat antes de aplicarla" % g("id"))
            continue
        if not g("operacion"):
            sin_op.append(g("id"))
            continue
        op = json.loads(g("operacion"))
        tipo = op["op"]
        if tipo not in OPS:
            print("%s: operación desconocida %s" % (g("id"), tipo))
            continue
        if a.aplicar:
            espera = None
            if tipo in validados and not a.lote:
                espera = "tipo validado: va con --lote"
            elif tipo not in validados and a.lote:
                espera = "el tipo %s no está validado; corre sin --lote para probarlo en una sola fila" % tipo
            elif tipo not in validados and (tipo in hechos_tipo or tipo in en_prueba):
                espera = "espera la validación de la primera fila de %s" % tipo
            if espera:
                try:  # still revalidate now, so a stale row is flagged before the batch
                    OPS[tipo](r, None, op, P, simula=True)
                    pendientes += 1
                    print("%s en espera: %s" % (g("id"), espera))
                except Revalidar as e:
                    print("%s REVALIDAR: %s" % (g("id"), e))
                    ws.cell(fila, COL["estado"]).value = "Revalidar"
                    ws.cell(fila, COL["comentario"]).value = ((g("comentario") or "") + " | Revalidar: %s" % e).strip(" |")
                continue
        if w:
            w.aprobacion = g("id")
            if tipo not in validados:
                hechos_tipo.add(tipo)          # one attempt per unvalidated type, even if it fails
                en_prueba.add(tipo)
                json.dump(sorted(en_prueba), open(prueba_path, "w"))
        try:
            msg, res = OPS[tipo](r, w, op, P, simula=not a.aplicar)
        except Revalidar as e:
            print("%s REVALIDAR: %s" % (g("id"), e))
            if a.aplicar:
                ws.cell(fila, COL["estado"]).value = "Revalidar"
                ws.cell(fila, COL["comentario"]).value = ((g("comentario") or "") + " | Revalidar: %s" % e).strip(" |")
            continue
        except Exception as e:
            print("%s ERROR: %s" % (g("id"), e))
            if a.aplicar:
                ws.cell(fila, COL["estado"]).value = "Error"
            continue
        print("%s %s" % (g("id"), msg))
        if a.aplicar and res:
            ws.cell(fila, COL["estado"]).value = "Aplicado"
            ws.cell(fila, COL["ids"]).value = ", ".join(str(i) for i in res["ids"])
            ws.cell(fila, COL["verificado"]).value = "Sí" if res["verificado"] else "No"
            hechos_tipo.add(tipo)
    if sin_op:
        print("Aprobadas sin detalle de operación (se completan en el chat): %s" % ", ".join(sin_op))
    if a.aplicar:
        wb.save(a.libro)
        nuevos = (hechos_tipo | en_prueba) - validados
        if nuevos:
            print("Primera aplicación de %s: revisa el resultado en Odoo con captura y, si está bien, corre --validar-tipo." % ", ".join(sorted(nuevos)))
        if pendientes:
            print("%d filas esperan la validación de su tipo para ir en lote." % pendientes)
    else:
        print("Simulación: no se escribió nada. Repite con --aplicar para la primera fila de cada tipo.")


if __name__ == "__main__":
    main()
