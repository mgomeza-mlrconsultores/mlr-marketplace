# -*- coding: utf-8 -*-
"""Builds a BBVA-like statement PDF from the demo workbook and checks that leer_estado_cuenta.py reads it
back with the control at zero. It proves the column-position reader, not BBVA's exact layout: the first
real statement of every new bank or format must still close its own control."""
import json, os, subprocess, sys, tempfile
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

AQUI = os.path.dirname(os.path.abspath(__file__))
EJ = os.path.join(AQUI, "..", "assets", "Ejemplo_Conciliacion_Agosto_2026_Demo.xlsx")
MES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
fmt = lambda v: "{:,.2f}".format(v)
ws = openpyxl.load_workbook(EJ)["Estado de cuenta"]
tmp = tempfile.mkdtemp()
pdf = os.path.join(tmp, "estado.pdf")
c = canvas.Canvas(pdf, pagesize=letter)
c.setFont("Helvetica", 8)
L = lambda r: ws.cell(r, 12).value
y = 740
for t in ["BBVA  Maestra PyME BBVA", "No. de Cuenta 0100000001   No. de Cliente E0000001", "CLABE 012180001000000011",
          "Saldo de Liquidación Inicial %s" % fmt(L(16)), "Depósitos / Abonos (+) %d %s" % (ws["M17"].value, fmt(L(17))),
          "Retiros / Cargos (-) %d %s" % (ws["M18"].value, fmt(L(18))), "Saldo Final (+) %s" % fmt(L(19))]:
    c.drawString(40, y, t); y -= 12
X = {"f1": 30, "f2": 62, "cod": 94, "desc": 118, "CARGOS": 400, "ABONOS": 460, "OPERACIÓN": 515, "LIQUIDACIÓN": 580}

def cabecera(y):
    c.drawString(X["f1"], y, "OPER"); c.drawString(X["f2"], y, "LIQ"); c.drawString(X["cod"], y, "COD.")
    c.drawString(X["desc"], y, "DESCRIPCIÓN")
    for k in ("CARGOS", "ABONOS", "OPERACIÓN", "LIQUIDACIÓN"):
        c.drawRightString(X[k] + 30, y, k)
y -= 10
cabecera(y); y -= 12
saldo = L(16)
for r in range(8, ws.max_row + 1):
    if not ws.cell(r, 1).value:
        continue
    if y < 60:
        c.showPage(); c.setFont("Helvetica", 8); y = 740; cabecera(y); y -= 12
    f = ws.cell(r, 2).value
    ff = "%02d/%s" % (f.day, MES[f.month - 1])
    cg, ab = ws.cell(r, 5).value, ws.cell(r, 6).value
    saldo = round(saldo + (ab or 0) - (cg or 0), 2)
    desc = ws.cell(r, 4).value
    partes = desc.split(" Ref. ")
    c.drawString(X["f1"], y, ff); c.drawString(X["f2"], y, ff); c.drawString(X["cod"], y, ws.cell(r, 3).value)
    c.drawString(X["desc"], y, partes[0][:48])
    if cg: c.drawRightString(X["CARGOS"] + 30, y, fmt(cg))
    if ab: c.drawRightString(X["ABONOS"] + 30, y, fmt(ab))
    c.drawRightString(X["OPERACIÓN"] + 30, y, fmt(saldo)); c.drawRightString(X["LIQUIDACIÓN"] + 30, y, fmt(saldo))
    y -= 11
    if len(partes) > 1:
        c.drawString(X["desc"], y, "Ref. " + partes[1][:60]); y -= 11
c.drawString(40, y - 10, "Total de Movimientos")
c.save()
out = os.path.join(tmp, "banco.json")
p = subprocess.run([sys.executable, os.path.join(AQUI, "..", "scripts", "leer_estado_cuenta.py"), pdf, out, "--anio", "2026"], capture_output=True, text=True)
print(p.stdout.strip())
d = json.load(open(out))
assert d["control"]["cuadra"], d["control"]
print("PDF sintético leído y cuadrado: %d movimientos" % len(d["movimientos"]))
