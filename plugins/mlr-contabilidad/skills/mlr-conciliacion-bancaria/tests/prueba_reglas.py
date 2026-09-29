# -*- coding: utf-8 -*-
"""Unit cases of the engine that the August example does not cover:
card-terminal lot on the same day, a USD journal, an invoice already paid before the period, own-account
transfers, and approvals that must not survive a changed operation."""
import json
import os
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, "..", "scripts"))
import openpyxl  # noqa: E402
from conciliar import Motor  # noqa: E402


def base(moneda="MXN"):
    P = {"rfc_empresa": "CDT260505AB1", "tolerancia_importe": 1.0, "tolerancia_dias": 3, "cuentas": {}, "diario_id": 1}
    odoo = {"compania": {"rfc": "CDT260505AB1"}, "diario": {"tipo": "bank", "cuenta": "102.01.01", "moneda": moneda},
            "periodo": {"desde": "2026-08-01", "hasta": "2026-08-31"}, "auxiliar": [], "facturas": {}, "iva": [], "diarios_caja": []}
    return P, odoo


def mov(p, fecha, desc, abono=None, cargo=None, codigo=""):
    return {"partida": p, "fecha": fecha, "codigo": codigo, "descripcion": desc, "abono": abono, "cargo": cargo, "referencia": "", "rfc": ""}


def aux(r, fecha, abono=None, cargo=None, contacto="", facturas=(), contrapartes=()):
    return {"renglon": r, "fecha": fecha, "asiento": "BNK/" + r, "concepto": "BNK/" + r, "contacto": contacto, "abono": abono, "cargo": cargo,
            "facturas": list(facturas), "contrapartes": list(contrapartes)}


def cfdi(u, total, receptor, fecha="2026-08-05", forma="04", moneda="MXN", tc=1.0, direccion="emitido"):
    return {"uuid": u, "direccion": direccion, "tipo": "I", "estado_sat": "Vigente", "fecha": fecha, "serie": "F", "folio": u[-3:],
            "total": total, "subtotal": round(total / 1.16, 2), "iva": round(total - total / 1.16, 2), "iva_ret": 0, "isr_ret": 0,
            "metodo_pago": "PUE", "forma_pago": forma, "moneda": moneda, "tipo_cambio": tc, "nombre_receptor": receptor,
            "rfc_receptor": "XAXX010101000", "fecha_pago_real": fecha, "pagos": []}


def corre(P, odoo, movs, cf, dec=None):
    banco = {"movimientos": movs, "caratula": {}, "control": {"cuadra": True}}
    M = Motor(P, banco, odoo, {"cfdi": {c["uuid"]: c for c in cf}, "indices": {"rep_por_factura": {}}}, dec)
    return {f["partida"]: f for f in M.filas() if f["primera"]}, M


def caso_lote_mismo_dia():
    P, O = base()
    O["auxiliar"] = [aux("O-001", "2026-08-05", abono=3000.0, contacto="Varios")]
    cf = [cfdi("A" * 33 + "001", 1000.0, "ANA"), cfdi("A" * 33 + "002", 1200.0, "LUIS"), cfdi("A" * 33 + "003", 800.0, "SARA"),
          cfdi("A" * 33 + "004", 5000.0, "OTRO", forma="03")]
    r, _ = corre(P, O, [mov("B-001", "2026-08-05", "VENTAS DEBITO TERMINALES", abono=3000.0, codigo="V42")], cf)
    assert r["B-001"]["enlace"] == "Varios CFDI" and r["B-001"]["regla"] == "E08", r["B-001"]


def caso_usd():
    P, O = base("USD")
    O["auxiliar"] = [aux("O-001", "2026-08-10", abono=100.0, contacto="CLIENTE USA")]
    cf = [cfdi("B" * 33 + "001", 100.0, "CLIENTE USA", fecha="2026-08-09", forma="03", moneda="USD", tc=18.5)]
    r, _ = corre(P, O, [mov("B-001", "2026-08-10", "SPEI RECIBIDO CLIENTE USA", abono=100.0)], cf)
    f = r["B-001"]
    assert f["enlace"] == "1 a 1" and f["total_cfdi"] == 100.0 and f["base_cfdi"] == round(round(100 / 1.16, 2) * 18.5, 2), f


def caso_pagada_antes():
    P, O = base()
    u = "C" * 33 + "001"
    O["facturas"] = {"7": {"id": 7, "numero": "F7", "uuid": u, "estado": "posted", "pendiente": 0.0, "total": 500.0, "partner": "ANA",
                           "tipo": "out_invoice", "del_mes": False}}
    O["auxiliar"] = [aux("O-001", "2026-08-10", abono=500.0, contacto="ANA")]
    r, _ = corre(P, O, [mov("B-001", "2026-08-10", "SPEI RECIBIDO ANA", abono=500.0)], [cfdi(u, 500.0, "ANA", fecha="2026-07-20", forma="03")])
    assert r["B-001"]["enlace"] == "Sin CFDI", r["B-001"]


def caso_traspaso_y_decision():
    P, O = base()
    O["auxiliar"] = [aux("O-001", "2026-08-10", cargo=20000.0, contrapartes=[{"codigo": "102.02.01", "tipo": "asset_cash"}]),
                     aux("O-002", "2026-08-11", cargo=3300.0, contacto="SOCIO")]
    r, _ = corre(P, O, [mov("B-001", "2026-08-10", "TRASPASO A INVERSION", cargo=20000.0), mov("B-002", "2026-08-11", "REEMBOLSO", cargo=3300.0)],
                 [], {"no_requiere": {"B-002": "Reembolso a socio"}})
    assert r["B-001"]["enlace"] == "No requiere CFDI" and r["B-001"]["estatus"] == "Conciliado", r["B-001"]
    assert r["B-002"]["enlace"] == "No requiere CFDI" and "Reembolso a socio" in r["B-002"]["observacion"], r["B-002"]


def caso_aprobacion_con_operacion_cambiada():
    from llenar_libro import conserva_capturas
    tmp = tempfile.mkdtemp()

    def libro(op, aprobado):
        wb = openpyxl.Workbook()
        for n in ("Previo", "Reglas"):
            wb.create_sheet(n)
        ws = wb.create_sheet("Acciones")
        for j, v in enumerate(["A001", "Registrar pago", "B-001", 10, "SPEI", "", "", "R01", "Alta", json.dumps(op), aprobado, "", "Pendiente", "", ""], 1):
            ws.cell(8, j).value = v
        return wb
    viejo = os.path.join(tmp, "viejo.xlsx")
    libro({"op": "registrar_pago", "factura_ids": [1007], "importe": 10}, "Sí").save(viejo)
    nuevo = libro({"op": "registrar_pago", "factura_ids": [99999], "importe": 10}, None)
    conserva_capturas(nuevo, viejo)
    ws = nuevo["Acciones"]
    assert ws.cell(8, 11).value in (None, "") and "vuelve a aprobar" in (ws.cell(8, 12).value or ""), (ws.cell(8, 11).value, ws.cell(8, 12).value)
    igual = libro({"op": "registrar_pago", "factura_ids": [1007], "importe": 10}, None)
    conserva_capturas(igual, viejo)
    assert igual["Acciones"].cell(8, 11).value == "Sí"


def caso_con_estado():
    P, O = base()
    P.update({"modo_diario": "con_estado", "cuentas": {"comisiones": "701.10.01"}})
    a = aux("O-001", "2026-08-03", cargo=299.7, contacto="")
    a["statement_line_id"] = 55
    O["auxiliar"] = [a]
    O["lineas_extracto_sin_conciliar"] = [{"id": 55, "fecha": "2026-08-03", "importe": -299.7}]
    r, M = corre(P, O, [mov("B-001", "2026-08-03", "APLI TASA DE DES", cargo=299.7, codigo="V46")], [])
    assert "suspenso" in r["B-001"]["observacion"], r["B-001"]
    acc = M.acciones(list(r.values()), [])
    op = next(x["operacion"] for x in acc if x["tipo"] == "Línea de extracto")
    assert op["op"] == "linea_extracto" and op["statement_line_id"] == 55 and op["contrapartidas"][0]["cuenta"] == "701.10.01", op


for caso in (caso_con_estado, caso_lote_mismo_dia, caso_usd, caso_pagada_antes, caso_traspaso_y_decision, caso_aprobacion_con_operacion_cambiada):
    caso()
    print("OK", caso.__name__)
print("Reglas: 6 casos correctos")
