# -*- coding: utf-8 -*-
"""Offline test of aplicar_acciones.py: a fake Odoo answers the RPC calls, so the whole write path runs
(backup, log, read-back, one-row-first gating, Revalidar) without touching a real database.

    python3 tests/prueba_escritura_simulada.py
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, "..", "scripts"))
import odoo_client  # noqa: E402

DB = {
    "account.account": {1: {"id": 1, "code": "701.10.01", "active": True}, 2: {"id": 2, "code": "102.01.01", "active": True}},
    "res.company": {1: {"id": 1, "fiscalyear_lock_date": False, "tax_lock_date": "2026-06-30", "hard_lock_date": False}},
    "account.move": {},
    "account.move.line": {10: {"id": 10, "reconciled": False, "account_id": [5, "105"], "partner_id": [9, "X"], "amount_residual": 100.0},
                          11: {"id": 11, "reconciled": False, "account_id": [5, "105"], "partner_id": [9, "X"], "amount_residual": -100.0}},
    "res.partner": {9: {"id": 9, "zip": "01000", "vat": "XAXX010101000"}},
}
LLAMADAS = []


def falso_rpc(self, service, method, args):
    if service == "common":
        return 7 if method == "authenticate" else {"server_version": "saas~19.4+e"}
    db, uid, key, model, meth, a, kw = args
    LLAMADAS.append((model, meth))
    tabla = DB.setdefault(model, {})
    if meth == "search_read":
        dom = a[0]
        out = []
        for rec in tabla.values():
            if all(rec.get(f) == v for f, op, v in [d for d in dom if isinstance(d, (list, tuple))] if op == "="):
                out.append(rec)
        return out
    if meth == "search_count":
        return 0
    if meth == "read":
        return [dict(tabla[i]) for i in a[0] if i in tabla]
    if meth == "create":
        nid = max(tabla or {0: 0}) + 1
        tabla[nid] = dict(a[0], id=nid, name="MISC/%d" % nid, state="draft", amount_total=sum(l[2]["debit"] for l in a[0].get("line_ids", [])))
        return nid
    if meth == "action_post":
        for i in a[0]:
            tabla[i]["state"] = "posted"
        return None
    if meth == "reconcile":
        for i in a[0]:
            tabla[i]["reconciled"], tabla[i]["amount_residual"] = True, 0.0
        return None
    if meth == "write":
        for i in a[0]:
            tabla[i].update(a[1])
        return True
    raise AssertionError("método inesperado %s.%s" % (model, meth))


def main():
    odoo_client.Lectura._rpc = falso_rpc
    os.environ.update({"ODOO_URL": "https://falso", "ODOO_DB": "falso", "ODOO_USER": "u@x", "ODOO_KEY": "a" * 40})
    tmp = tempfile.mkdtemp()
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acciones"
    ops = [{"op": "asiento_manual", "diario_id": 3, "fecha": "2026-08-31", "referencia": "Comisión", "lineas": [{"cuenta": "701.10.01", "debe": 10}, {"cuenta": "102.01.01", "haber": 10}]},
           {"op": "asiento_manual", "diario_id": 3, "fecha": "2026-08-31", "referencia": "Comisión 2", "lineas": [{"cuenta": "701.10.01", "debe": 5}, {"cuenta": "102.01.01", "haber": 5}]},
           {"op": "asiento_manual", "diario_id": 3, "fecha": "2026-06-15", "referencia": "Fecha bloqueada", "lineas": [{"cuenta": "701.10.01", "debe": 5}, {"cuenta": "102.01.01", "haber": 5}]},
           {"op": "conciliar_apuntes", "line_ids": [10, 11]},
           {"op": "escribir_campos", "modelo": "res.partner", "ids": [9], "valores": {"zip": "06600"}}]
    for i, op in enumerate(ops, 8):
        for j, v in enumerate(["T%d" % i, "x", "", 0, "", "", "", "", "", json.dumps(op), "Sí", "", "Pendiente", "", ""], 1):
            ws.cell(i, j).value = v
    libro = os.path.join(tmp, "libro.xlsx")
    wb.save(libro)
    params = os.path.join(tmp, "params.json")
    json.dump({"compania_id": 1, "diario_id": 3, "entorno": "pruebas"}, open(params, "w"))
    import aplicar_acciones as AA
    json.dump({"compania_id": 1, "diario_id": 3}, open(params, "w"))  # no entorno declared: must refuse
    sys.argv = ["aplicar_acciones.py", params, libro, tmp, "--aplicar"]
    try:
        AA.main()
        raise AssertionError("escribió sin entorno declarado")
    except SystemExit as e:
        assert "entorno" in str(e)
    json.dump({"compania_id": 1, "diario_id": 3, "entorno": "produccion"}, open(params, "w"))  # production without today's backup
    try:
        AA.main()
        raise AssertionError("escribió en producción sin respaldo")
    except SystemExit as e:
        assert "respaldo_base" in str(e)
    json.dump({"compania_id": 1, "diario_id": 3, "entorno": "pruebas"}, open(params, "w"))
    def corre(*extra):
        sys.argv = ["aplicar_acciones.py", params, libro, tmp] + list(extra)
        AA.main()
    corre("--aplicar")                       # first of each type only
    est = [openpyxl.load_workbook(libro)["Acciones"].cell(r, 13).value for r in range(8, 13)]
    assert est == ["Aplicado", "Pendiente", "Revalidar", "Aplicado", "Aplicado"], est
    corre("--aplicar")                       # a second run must not apply more of an unvalidated type
    assert openpyxl.load_workbook(libro)["Acciones"].cell(9, 13).value == "Pendiente"
    corre("--aplicar", "--lote")             # asiento_manual not validated yet: nothing more
    assert openpyxl.load_workbook(libro)["Acciones"].cell(9, 13).value == "Pendiente"
    corre("--validar-tipo", "asiento_manual")
    corre("--aplicar", "--lote")
    wsr = openpyxl.load_workbook(libro)["Acciones"]
    assert wsr.cell(9, 13).value == "Aplicado" and wsr.cell(9, 15).value == "Sí"
    bit = [json.loads(x) for x in open(os.path.join(tmp, "bitacora.jsonl"))]
    assert all(os.path.exists(b["respaldo"]) for b in bit)
    assert DB["res.partner"][9]["zip"] == "06600"
    print("Escritura simulada: %d escrituras en bitácora, prueba antes de lote y revalidación correctas" % len(bit))
    shutil.rmtree(tmp)


if __name__ == "__main__":
    main()
